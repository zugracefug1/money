"""엑셀·CSV 거래내역 읽기 → 헤더/열 감지 → 값 정규화 → 거래 행 만들기.

원본 파일은 메모리에서만 읽고 어디에도 저장하지 않는다. 카드번호·계좌번호로 보이는 열은 매핑에서 뺀다.
"""

from __future__ import annotations

import calendar
import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any

Cell = Any  # str | int | float | bool | datetime | date | time | None

ACCEPTED_EXTENSIONS = (".xlsx", ".xls", ".csv")
MAX_FILE_BYTES = 10 * 1024 * 1024

FIELDS = ("date", "time", "merchant", "amount", "memo")
FIELD_LABELS = {"date": "날짜", "time": "시간", "merchant": "사용처", "amount": "금액", "memo": "메모"}
REQUIRED_FIELDS = ("date", "merchant", "amount")


# 읽기 ---------------------------------------------------------------------------


@dataclass
class RawTable:
    sheet_name: str
    rows: list[list[Cell]]
    first_row_number: int  # rows[0] 이 엑셀에서 몇 번째 행인지(1부터)


def extension_of(file_name: str) -> str:
    dot = file_name.rfind(".")
    return "" if dot < 0 else file_name[dot:].lower()


def decode_csv(data: bytes) -> str:
    """국내 은행 CSV는 EUC-KR(CP949)인 경우가 많다. UTF-8로 읽어 보고 실패하면 CP949로 읽는다."""
    try:
        return data.decode("utf-8").lstrip("﻿")
    except UnicodeDecodeError:
        return data.decode("cp949", errors="replace")


def _is_filled(c: Cell) -> bool:
    return c is not None and str(c).strip() != ""


def _clean_cell(v: Cell) -> Cell:
    if isinstance(v, float) and v != v:  # NaN
        return None
    return v


def read_workbook(data: bytes, file_name: str) -> RawTable:
    """시트가 여러 개면 데이터가 가장 많은 시트를 고른다."""
    ext = extension_of(file_name)
    tables: list[RawTable] = []

    if ext == ".csv":
        text = decode_csv(data)
        rows = [[(c if c != "" else None) for c in r] for r in csv.reader(io.StringIO(text))]
        tables.append(RawTable("CSV", rows, 1))
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [[_clean_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
            tables.append(RawTable(ws.title, rows, ws.min_row or 1))
        wb.close()
    elif ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            tables.append(RawTable(sh.name, rows, 1))
    else:
        raise ValueError("xlsx, xls, csv 파일만 올릴 수 있습니다.")

    best, best_count = None, 0
    for t in tables:
        count = sum(1 for r in t.rows if any(_is_filled(c) for c in r))
        if count > best_count:
            best, best_count = t, count
    if best is None:
        raise ValueError("파일에서 데이터를 찾을 수 없습니다.")
    return best


# 헤더·열 감지 ---------------------------------------------------------------------

# 앞에 있는 단어일수록 우선한다
KEYWORDS = {
    "date": ["이용일자", "이용일시", "거래일자", "거래일시", "승인일자", "승인일시", "사용일자", "이용일", "거래일", "승인일", "사용일", "날짜", "일자", "일시"],
    "time": ["이용시간", "승인시간", "거래시간", "사용시간", "시간", "시각"],
    "merchant": ["가맹점명", "가맹점", "이용하신곳", "이용처", "사용처", "상호", "거래처", "이용내역", "적요", "내용"],
    "amount": ["이용금액", "승인금액", "거래금액", "사용금액", "출금액", "출금금액", "결제금액", "결제원금", "금액", "출금"],
    "memo": ["메모", "비고", "기재내용", "내용"],
}
# 금액 열로 고르면 안 되는 열(해외 통화, 수수료, 잔액, 입금 등)
AMOUNT_EXCLUDE = ["해외", "달러", "USD", "외화", "수수료", "할인", "포인트", "적립", "입금", "잔액", "이자", "할부"]

SENSITIVE_HEADER = re.compile(r"카드\s*번호|카드\s*No|계좌|이용카드|주민|card\s*no|account", re.I)
SENSITIVE_VALUE = re.compile(r"\d{4}[-\s]?[\d*xX]{4}[-\s]?[\d*xX]{4}[-\s]?[\d*xX]{3,4}|\d{4}-\*{4}")


def normalize_header(value: Cell) -> str:
    return "" if value is None else re.sub(r"\s+", "", str(value)).strip()


def _keyword_score(fld: str, header: str) -> int:
    if not header:
        return 0
    if fld == "amount" and any(w in header for w in AMOUNT_EXCLUDE):
        return 0
    for i, kw in enumerate(KEYWORDS[fld]):
        if header == kw:
            return 200 - i
        if kw in header:
            return 100 - i
    return 0


def find_header_row(rows: list[list[Cell]]) -> int:
    """머리말(제목·조회기간 등) 아래의 실제 헤더 행: 앞쪽 30행 중 관련 단어가 가장 많이 나오는 행."""
    best, best_score = -1, 0
    for index, row in enumerate(rows[:30]):
        headers = [normalize_header(c) for c in row]
        matched = sum(1 for f in ("date", "merchant", "amount", "time", "memo") if any(_keyword_score(f, h) > 0 for h in headers))
        score = matched * 100 + sum(1 for c in row if _is_filled(c))
        if matched >= 2 and score > best_score:
            best, best_score = index, score
    if best >= 0:
        return best
    for index, row in enumerate(rows):
        if sum(1 for c in row if _is_filled(c)) >= 2:
            return index
    return 0


def find_sensitive_columns(header_row: list[Cell], data_rows: list[list[Cell]]) -> set[int]:
    """카드번호·계좌번호처럼 보이는 열 번호들. 헤더 이름이나 아래 값의 모양으로 판단한다."""
    sensitive = {col for col, h in enumerate(header_row) if SENSITIVE_HEADER.search(normalize_header(h))}
    for row in data_rows[:50]:
        for col, v in enumerate(row):
            if isinstance(v, str) and SENSITIVE_VALUE.search(v):
                sensitive.add(col)
    return sensitive


def guess_mapping(header_row: list[Cell], exclude: set[int] | None = None) -> dict[str, int]:
    """헤더 이름으로 필드별 열을 추측한다. 한 열은 한 필드에만 쓴다."""
    headers = [normalize_header(c) for c in header_row]
    mapping: dict[str, int] = {}
    used = set(exclude or ())
    # 날짜·금액·사용처를 먼저 정해 '내용' 같은 겹치는 단어가 메모로 먼저 가지 않게 한다
    for fld in ("date", "amount", "merchant", "time", "memo"):
        best_col, best_score = -1, 0
        for col, h in enumerate(headers):
            if col in used:
                continue
            score = _keyword_score(fld, h)
            if score > best_score:
                best_col, best_score = col, score
        if best_col >= 0:
            mapping[fld] = best_col
            used.add(best_col)
    return mapping


def header_signature(header_row: list[Cell]) -> str:
    return "|".join(h for h in (normalize_header(c) for c in header_row) if h)


def to_preset_mapping(header_row: list[Cell], mapping: dict[str, int]) -> dict:
    """열 번호 대신 헤더 이름을 저장해 열 순서가 바뀌어도 쓸 수 있게 한다."""
    columns = {f: normalize_header(header_row[c]) for f, c in mapping.items() if c is not None}
    return {"signature": header_signature(header_row), "columns": columns}


def from_preset_mapping(header_row: list[Cell], preset: dict, exclude: set[int] | None = None) -> dict[str, int] | None:
    headers = [normalize_header(c) for c in header_row]
    exclude = exclude or set()
    mapping: dict[str, int] = {}
    for fld in FIELDS:
        name = (preset.get("columns") or {}).get(fld)
        if not name or name not in headers:
            continue
        col = headers.index(name)
        if col not in exclude:
            mapping[fld] = col
    return mapping if all(f in mapping for f in REQUIRED_FIELDS) else None


def is_preset_mapping(value: Any) -> bool:
    return isinstance(value, dict) and isinstance(value.get("signature"), str) and isinstance(value.get("columns"), dict)


# 값 정규화 -------------------------------------------------------------------------


def _valid_date(y: int, m: int, d: int) -> str | None:
    if m < 1 or m > 12 or d < 1 or y < 1990 or y > 2100:
        return None
    if d > calendar.monthrange(y, m)[1]:
        return None
    return f"{y:04d}-{m:02d}-{d:02d}"


_EXCEL_EPOCH = datetime(1899, 12, 30)


def _from_excel_serial(serial: float) -> tuple[str | None, str | None]:
    dt = _EXCEL_EPOCH + timedelta(milliseconds=round(serial * 86400 * 1000))
    d = _valid_date(dt.year, dt.month, dt.day)
    t = dt.strftime("%H:%M:%S") if serial % 1 != 0 else None
    return d, t


TIME_RE = re.compile(r"(오전|오후|AM|PM)?\s*(\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM)?", re.I)


def _parse_time_text(text: str) -> str | None:
    m = TIME_RE.search(text)
    if not m:
        return None
    h, mi, s = int(m.group(2)), int(m.group(3)), int(m.group(4) or 0)
    meridiem = (m.group(1) or m.group(5) or "").upper()
    if meridiem in ("오후", "PM"):
        h = h % 12 + 12
    if meridiem in ("오전", "AM"):
        h = h % 12
    if h > 23 or mi > 59 or s > 59:
        return None
    return f"{h:02d}:{mi:02d}:{s:02d}"


DATE_PATTERNS = [
    # 2026.10.01 · 2026-10-01 · 2026/10/01 · 2026년 10월 1일
    (re.compile(r"^(\d{4})\s*[.\-/년]\s*(\d{1,2})\s*[.\-/월]\s*(\d{1,2})\s*일?\.?"), lambda s: int(s)),
    # 20261001
    (re.compile(r"^(\d{4})(\d{2})(\d{2})(?!\d)"), lambda s: int(s)),
    # 26.10.01 · 26/10/01 · 26-10-01
    (re.compile(r"^(\d{2})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*(\d{1,2})(?!\d)"), lambda s: 2000 + int(s)),
]


def parse_date(value: Cell) -> tuple[str, str | None] | None:
    """날짜 셀 → (YYYY-MM-DD, HH:MM:SS | None). 같은 칸에 시간이 붙어 있으면 시간도 돌려준다."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        d = _valid_date(value.year, value.month, value.day)
        t = value.strftime("%H:%M:%S") if (value.hour, value.minute, value.second) != (0, 0, 0) else None
        return (d, t) if d else None
    if isinstance(value, date):
        d = _valid_date(value.year, value.month, value.day)
        return (d, None) if d else None
    if isinstance(value, (int, float)):
        # 1990~2100 범위의 엑셀 일련번호만 날짜로 본다
        if value < 32874 or value > 73051:
            return None
        d, t = _from_excel_serial(float(value))
        return (d, t) if d else None
    text = str(value).strip()
    for pattern, year in DATE_PATTERNS:
        m = pattern.match(text)
        if not m:
            continue
        d = _valid_date(year(m.group(1)), int(m.group(2)), int(m.group(3)))
        if not d:
            return None
        return d, _parse_time_text(text[m.end():])
    return None


def parse_time(value: Cell) -> str | None:
    """시간 셀 → HH:MM:SS. 엑셀 시간 값(하루의 비율)과 "14:05", "오후 2:05", "140530" 형식을 받는다."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (int, float)):
        fraction = value % 1
        if value < 0 or (value >= 1 and fraction == 0):
            return None
        return _from_excel_serial(25569 + fraction)[1] or "00:00:00"
    text = str(value).strip()
    compact = re.fullmatch(r"(\d{2})(\d{2})(\d{2})?", text)
    if compact:
        return _parse_time_text(f"{compact.group(1)}:{compact.group(2)}:{compact.group(3) or '00'}")
    return _parse_time_text(text)


_MISSING = object()


def parse_amount(value: Cell):
    """금액 셀 → 원 단위 정수. "12,000원", "₩12,000", "-3,500", "(3,500)", "△3,500"(음수)를 받는다.

    빈 칸은 _MISSING(→ "금액 없음"), 해석할 수 없으면 None.
    """
    if value is None:
        return _MISSING
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        return int(round(value)) if value >= 0 else -int(round(-value))
    text = re.sub(r"\s+", "", str(value).strip())
    if text in ("", "-"):
        return _MISSING
    negative = False
    if re.fullmatch(r"\(.*\)", text):
        negative, text = True, text[1:-1]
    if re.match(r"^[-−△▲]", text):
        negative, text = True, text[1:]
    text = re.sub(r"^(₩|￦|KRW)", "", text, flags=re.I)
    text = re.sub(r"(원|KRW)$", "", text, flags=re.I).replace(",", "")
    if not re.fullmatch(r"\d+(\.\d+)?", text):
        return None
    amount = int(float(text) + 0.5)
    return -amount if negative else amount


def is_missing(v) -> bool:
    return v is _MISSING


def clean_text(value: Cell) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


# 행 만들기 ---------------------------------------------------------------------------

SUMMARY_RE = re.compile(r"^(합계|소계|총계|총\s*합계|합\s*계|total)", re.I)


@dataclass
class RowResult:
    row_number: int
    status: str  # "ok" | "error"
    tx: dict | None = None  # date, time, merchant, amount, memo
    reason: str = ""
    cells: list = field(default_factory=list)
    duplicate: bool = False


def _blank(c: Cell) -> bool:
    return c is None or str(c).strip() == ""


def build_rows(table: RawTable, header_index: int, mapping: dict[str, int]) -> tuple[list[RowResult], int]:
    """헤더 아래 행들을 거래로 바꾼다. 실패한 행은 이유와 함께 error 로 남긴다. (결과, 건너뛴 행 수)"""
    results: list[RowResult] = []
    skipped = 0
    if any(f not in mapping for f in REQUIRED_FIELDS):
        return results, skipped

    def pick(cells, col):
        return cells[col] if col is not None and col < len(cells) else None

    for i, cells in enumerate(table.rows[header_index + 1 :]):
        row_number = table.first_row_number + header_index + 1 + i
        raw_date = pick(cells, mapping["date"])
        if all(_blank(c) for c in cells):
            skipped += 1
            continue
        if any(isinstance(c, str) and SUMMARY_RE.match(c.strip()) for c in cells) and parse_date(raw_date) is None:
            skipped += 1
            continue

        parsed = parse_date(raw_date)
        merchant = clean_text(pick(cells, mapping["merchant"]))
        amount = parse_amount(pick(cells, mapping["amount"]))
        reason = (
            ("날짜 없음" if _blank(raw_date) else "날짜 형식 오류") if not parsed
            else "사용처 없음" if not merchant
            else "금액 없음" if is_missing(amount)
            else "금액 형식 오류" if amount is None
            else "금액 0원" if amount == 0
            else None
        )  # fmt: skip
        if reason:
            results.append(RowResult(row_number, "error", reason=reason, cells=list(cells)))
            continue
        t = parse_time(pick(cells, mapping["time"])) if "time" in mapping else parsed[1]
        memo = clean_text(pick(cells, mapping["memo"])) or None if "memo" in mapping else None
        results.append(
            RowResult(row_number, "ok", tx={"date": parsed[0], "time": t, "merchant": merchant, "amount": amount, "memo": memo})
        )
    return results, skipped


def _key(d: dict) -> tuple:
    return (d["date"], d["merchant"], d["amount"])


def mark_duplicates(results: list[RowResult], existing: list[dict]) -> list[RowResult]:
    """이미 저장된 거래와 같은 날짜·사용처·금액인 행을 중복으로 표시한다.

    같은 날 같은 곳에서 같은 금액을 두 번 쓴 경우를 살리려고 개수로 비교한다(기존 2건·파일 3건 → 앞의 2건만 중복).
    """
    remaining: dict[tuple, int] = {}
    for e in existing:
        remaining[_key(e)] = remaining.get(_key(e), 0) + 1
    for r in results:
        if r.status != "ok":
            continue
        left = remaining.get(_key(r.tx), 0)
        if left > 0:
            remaining[_key(r.tx)] = left - 1
        r.duplicate = left > 0
    return results
